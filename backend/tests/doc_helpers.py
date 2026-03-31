"""Build real PDF/DOCX files for document tests."""

import textwrap
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas


def make_pdf(path: Path, pages: list[str]) -> Path:
    """One PDF page per entry; an empty string makes a page with no text layer."""
    pdf = canvas.Canvas(str(path))
    for text in pages:
        y = 800
        for line in textwrap.wrap(text, 90):
            pdf.drawString(40, y, line)
            y -= 14
        pdf.showPage()
    pdf.save()
    return path


def make_encrypted_pdf(
    path: Path,
    text: str,
    password: str = "secret",
    algorithm: str = "RC4-128",
    owner_password: str | None = None,
) -> Path:
    """Encrypted PDF. With password="" it opens without a password (owner restrictions only)."""
    plain = make_pdf(path.with_suffix(".plain.pdf"), [text])
    writer = PdfWriter()
    for page in PdfReader(plain).pages:
        writer.add_page(page)
    writer.encrypt(
        user_password=password, owner_password=owner_password or "owner-secret", algorithm=algorithm
    )
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def make_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> Path:
    doc = DocxDocument()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    if table:
        grid = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    doc.save(str(path))
    return path
