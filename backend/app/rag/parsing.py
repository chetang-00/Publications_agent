"""Turn uploaded files into text, keeping page numbers where the format has them."""

import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

log = logging.getLogger(__name__)

DocKind = Literal["pdf", "docx", "txt", "md"]

EXTENSIONS: dict[str, DocKind] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
    ".md": "md",
    ".markdown": "md",
}
CONTENT_TYPES: dict[DocKind, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "txt": "text/plain",
    "md": "text/markdown",
}
KIND_BY_CONTENT_TYPE: dict[str, DocKind] = {v: k for k, v in CONTENT_TYPES.items()}
MIN_TEXT_CHARS = 20


class UnsupportedFile(Exception):
    """Rejected at upload time (HTTP 400)."""


class ParseError(Exception):
    """The file was accepted but no usable text could be extracted (document marked failed)."""


@dataclass
class ParsedPage:
    page: int | None  # 1-based page number; None for formats without pages
    text: str


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]  # pages with text only
    page_count: int


def detect_kind(filename: str, head: bytes) -> DocKind:
    suffix = Path(filename).suffix.lower()
    kind = EXTENSIONS.get(suffix)
    if kind is None:
        raise UnsupportedFile(
            f"Unsupported file type '{suffix or filename}'. Upload PDF, DOCX, TXT or MD files."
        )
    if not head:
        raise UnsupportedFile("The file is empty.")
    if kind == "pdf" and not head.startswith(b"%PDF-"):
        raise UnsupportedFile("The file has a .pdf name but is not a PDF document.")
    if kind == "docx" and not head.startswith(b"PK\x03\x04"):
        raise UnsupportedFile("The file has a .docx name but is not a Word document.")
    if kind in ("txt", "md") and b"\x00" in head:
        raise UnsupportedFile("The file looks binary, not plain text.")
    return kind


def _tidy(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0
            if not unlocked:
                raise ParseError("The PDF is password-protected. Remove the password and upload it again.")
        pages = [
            ParsedPage(page=i + 1, text=_tidy(page.extract_text() or ""))
            for i, page in enumerate(reader.pages)
        ]
    except ParseError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise ParseError(f"Could not read the PDF: {exc}") from exc
    except Exception as exc:  # pypdf raises a wide range of errors on malformed files
        log.warning("Unexpected PDF parse failure", exc_info=True)
        raise ParseError(f"Could not read the PDF ({type(exc).__name__}).") from exc
    return ParsedDocument(pages=[p for p in pages if p.text], page_count=len(pages))


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        doc = DocxDocument(str(path))
    except (zipfile.BadZipFile, KeyError, ValueError, PackageNotFoundError) as exc:
        raise ParseError("Could not read the Word document; the file is damaged or not a .docx.") from exc
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _tidy("\n\n".join(parts))
    return ParsedDocument(pages=[ParsedPage(page=None, text=text)] if text else [], page_count=1)


def _parse_text(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")
    text = _tidy(text)
    return ParsedDocument(pages=[ParsedPage(page=None, text=text)] if text else [], page_count=1)


def parse_document(path: Path, kind: DocKind) -> ParsedDocument:
    if kind == "pdf":
        parsed = _parse_pdf(path)
    elif kind == "docx":
        parsed = _parse_docx(path)
    else:
        parsed = _parse_text(path)
    if sum(len(p.text.strip()) for p in parsed.pages) < MIN_TEXT_CHARS:
        raise ParseError(
            "No extractable text found. Scanned or image-only PDFs are not supported (there is no OCR)."
        )
    return parsed
